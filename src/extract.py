"""F3 — Extracción multi-formato del corpus jurídico.

Lee los artefactos preservados por ``src.ingest`` y produce
``data/processed/extracted.jsonl`` con un registro por documento:

    {doc_id, titulo, fuente_url, formato, tipo_norma, materia, texto}

Soporta HTML, PDF, Word (.docx) y Excel (.xls/.xlsx). El texto se normaliza
pero **no** se fragmenta aquí (eso es responsabilidad de ``src.chunking``).
"""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from urllib.parse import urlparse

from config import settings

TAGS_A_ELIMINAR = ["script", "style", "nav", "header", "footer", "form",
                   "noscript", "iframe", "svg", "aside", "button"]
# Contenedores de "ruido" del portal (Drupal): navegación, menús, sidebars…
RE_CLASES_RUIDO = re.compile(
    r"(menu|nav|breadcrumb|sidebar|skip|region-header|region-footer|"
    r"region-navigation|region-services|region-sidebar|block-search|"
    r"panel|pager|tabs|social|banner|cookie|submenu|dhtml)",
    re.IGNORECASE,
)
# Selectores preferidos para el contenido útil, en orden de prioridad.
SELECTORES_CONTENIDO = ["main", "article", ".region-content", "#content",
                        ".view-content", ".main-container"]
BLOQUES = ["p", "div", "section", "article", "li", "tr", "br",
           "h1", "h2", "h3", "h4", "h5", "h6", "table", "td", "th"]

# Heurísticas para metadatos jurídicos a partir de URL/título.
REGLAS_TIPO = [
    (r"ley", "Ley"),
    (r"convenio", "Convenio cambiario"),
    (r"resoluci", "Resolución"),
    (r"circular", "Circular"),
    (r"aviso", "Aviso oficial"),
    (r"providencia", "Providencia"),
    (r"decreto", "Decreto"),
    (r"acto", "Acto administrativo"),
    (r"reglamento", "Reglamento"),
]
REGLAS_MATERIA = [
    (r"cambiar", "Cambiario"),
    (r"monetar", "Monetario"),
    (r"pago", "Sistema de pagos"),
    (r"mineral", "Minerales estratégicos"),
    (r"inpc|precio", "Precios e inflación"),
    (r"encaje", "Encaje legal"),
    (r"titulo", "Títulos de cobertura"),
]


def _norm(texto: str) -> str:
    lineas = []
    for ln in texto.split("\n"):
        ln = re.sub(r"[ \t\u00a0]+", " ", ln).strip()
        if ln:
            lineas.append(ln)
    return "\n".join(lineas)


def _sha(texto: str) -> str:
    return hashlib.sha256(texto.encode("utf-8")).hexdigest()


def _inferir(url: str, titulo: str) -> tuple[str, str]:
    blob = f"{url} {titulo}".lower()
    tipo = next((v for pat, v in REGLAS_TIPO if re.search(pat, blob)), "Documento normativo")
    materia = next((v for pat, v in REGLAS_MATERIA if re.search(pat, blob)), "General")
    return tipo, materia


# ----------------------------------------------------------------------
# Parsers
# ----------------------------------------------------------------------
def parse_html_file(path: Path) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "lxml")

    # 1) Etiquetas no textuales
    for tag in soup(TAGS_A_ELIMINAR):
        tag.decompose()

    # 2) Contenedores de ruido del portal (menús, sidebars, cabeceras…)
    for el in soup.find_all(True):
        if el.parent is None:
            continue
        ident = " ".join(el.get("class") or []) + " " + (el.get("id") or "")
        if ident.strip() and RE_CLASES_RUIDO.search(ident):
            el.decompose()

    # 3) Aislar el contenedor de contenido útil
    cont = None
    for sel in SELECTORES_CONTENIDO:
        cont = soup.select_one(sel)
        if cont is not None:
            break
    cont = cont or soup.body or soup

    for br in cont.find_all("br"):
        br.replace_with("\n")
    for b in cont.find_all(BLOQUES):
        b.insert_before("\n")
    return _norm(cont.get_text(separator=" "))


def parse_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    partes = [t for p in reader.pages if (t := p.extract_text())]
    return _norm("\n".join(partes))


def parse_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    partes = [p.text for p in d.paragraphs if p.text.strip()]
    for tabla in d.tables:
        for fila in tabla.rows:
            partes.append(" | ".join(c.text.strip() for c in fila.cells))
    return _norm("\n".join(partes))


def parse_excel(path: Path) -> str:
    import pandas as pd

    libro = pd.ExcelFile(path)
    bloques = []
    for hoja in libro.sheet_names:
        df = libro.parse(hoja, header=None)
        bloques.append(f"[Hoja: {hoja}]")
        bloques.append(df.to_csv(index=False, header=False))
    return _norm("\n".join(bloques))


PARSERS = {
    ".html": parse_html_file,
    ".htm": parse_html_file,
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_excel,
    ".xls": parse_excel,
}


def _extraer(path: Path) -> str:
    fn = PARSERS.get(path.suffix.lower())
    if fn is None:
        raise ValueError(f"Formato no soportado: {path.suffix}")
    return fn(path)


# ----------------------------------------------------------------------
# Pipeline
# ----------------------------------------------------------------------
def run() -> list[dict]:
    manifest_path = settings.RAW_DIR / "manifest.json"
    docs_path = settings.RAW_DIR / "documentos.json"
    if not manifest_path.exists():
        raise SystemExit("No hay manifest.json. Ejecuta primero: python -m src.ingest")

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    documentos = json.loads(docs_path.read_text(encoding="utf-8")) if docs_path.exists() else []

    registros: list[dict] = []
    vistos: set[str] = set()
    vacios = 0

    print("Extrayendo páginas HTML...\n")
    for p in manifest:
        if p.get("estado") != "ok":
            continue
        path = settings.RAW_DIR / p["archivo"]
        try:
            texto = _extraer(path)
        except Exception as exc:  # noqa: BLE001
            print(f"  ERROR {p['archivo']} -> {exc}")
            continue
        if len(texto) < 40:
            vacios += 1
            continue
        tipo, materia = _inferir(p["url"], p["titulo"])
        registros.append({
            "doc_id": p["id"], "titulo": p["titulo"], "fuente_url": p["url"],
            "formato": "html", "tipo_norma": tipo, "materia": materia,
            "hash": _sha(texto), "texto": texto, "n_caracteres": len(texto),
        })
        print(f"  OK    {p['id']:8s} {p['titulo'][:48]:<48} {len(texto):>7d} car.")

    print(f"\nExtrayendo {len(documentos)} documentos (PDF/Word/Excel)...\n")
    for d in documentos:
        if d.get("estado") != "ok":
            continue
        path = settings.RAW_DIR / "docs" / d["archivo"]
        if not path.exists():
            continue
        try:
            texto = _extraer(path)
        except Exception as exc:  # noqa: BLE001
            print(f"  ERROR {d['archivo']} -> {exc}")
            continue
        if len(texto) < 40:
            vacios += 1
            print(f"  VACIO {d['archivo']} (posible escaneado)")
            continue
        h = _sha(texto)
        if h in vistos:
            print(f"  DUP   {d['archivo']}")
            continue
        vistos.add(h)
        tipo, materia = _inferir(d["url"], d["archivo"])
        registros.append({
            "doc_id": d["id"], "titulo": d["archivo"], "fuente_url": d["url"],
            "formato": d["ext"].lstrip("."), "tipo_norma": tipo, "materia": materia,
            "hash": h, "texto": texto, "n_caracteres": len(texto),
        })
        print(f"  OK    {d['id']:8s} {d['archivo'][:48]:<48} {len(texto):>7d} car.")

    out = settings.PROCESSED_DIR / "extracted.jsonl"
    with out.open("w", encoding="utf-8") as f:
        for r in registros:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    print(f"\nResumen extracción: {len(registros)} documentos · {vacios} vacíos")
    print(f"Salida: {out}")
    return registros


if __name__ == "__main__":
    run()
